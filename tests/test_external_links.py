from pathlib import Path

from openpyxl import Workbook, load_workbook


def test_generated_workbook_has_no_external_links(tmp_path: Path):
    output_path = tmp_path / "test.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "テスト"
    workbook.save(output_path)
    workbook.close()

    loaded_workbook = load_workbook(output_path, data_only=False)

    try:
        assert loaded_workbook._external_links == []
    finally:
        loaded_workbook.close()