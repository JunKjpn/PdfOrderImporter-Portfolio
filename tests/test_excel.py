import pandas as pd
from openpyxl import Workbook

from src.excel.summary import orders_to_dataframe
from src.excel.writer import write_orders


def test_write_orders():
    workbook = Workbook()
    worksheet = workbook.active

    merged_df = pd.DataFrame(
        {
            "商品名": [
                "商品A",
                "商品B",
                "商品C(バリエーション1)",
            ],
            "数量": [10, 5, 8],
        }
    )

    product_rows = {
        "商品A": 5,
        "商品B": 6,
        "商品C(バリエーション1)": 7,
    }

    write_orders(
        worksheet=worksheet,
        merged_df=merged_df,
        start_row=5,
        quantity_column=3,
    )

    assert worksheet.cell(row=5, column=3).value == 10
    assert worksheet.cell(row=6, column=3).value == 5
    assert worksheet.cell(row=7, column=3).value == 8


def test_orders_to_dataframe():
    orders = {
        "商品A": 10,
        "商品B": 5,
    }

    result = orders_to_dataframe(orders)

    assert result.to_dict("records") == [
        {"商品名": "商品A", "数量": 10},
        {"商品名": "商品B", "数量": 5},
    ]