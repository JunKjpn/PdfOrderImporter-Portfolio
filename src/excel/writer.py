from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.excel.finder import find_column_by_value
from src.models.product import Product


def open_workbook(template_path: Path, sheet_name: str) -> tuple[Workbook, Worksheet]:
    """テンプレートExcelを開き、WorkbookとWorksheetを返す。"""
    workbook = load_workbook(template_path)
    worksheet = workbook[sheet_name]
    return workbook, worksheet

def save_workbook(workbook: Workbook, output_path: Path) -> None:
    """Excelを保存し、閉じる。"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

def clear_cell_range(worksheet: Worksheet, start_row: int, end_row: int, start_column: int, end_column: int) -> None:
    """指定した範囲のセル値をクリアする。"""
    for row_number in range(start_row, end_row + 1):
        for column_number in range(start_column, end_column + 1):
            worksheet.cell(row=row_number, column=column_number).value = None

def write_product_names(worksheet: Worksheet, merged_df: pd.DataFrame, start_row: int, product_column: int) -> None:
    """注文された商品名をExcelへ書き込む。"""
    for row, product_name in enumerate(merged_df["商品名"], start=start_row):
        worksheet.cell(row=row, column=product_column).value = product_name

def write_orders(worksheet: Worksheet, merged_df: pd.DataFrame, start_row: int, quantity_column: int) -> None:
    """商品マスタ順に整列された発注数量をExcelへ書き込む。"""
    for row, quantity in enumerate(merged_df["数量"], start=start_row):
        worksheet.cell(row=row, column=quantity_column).value = quantity

def write_product_master_value(
    worksheet: Worksheet,
    product_master: dict[str, Product],
    product_names: list[str],
    search_value: str,
    header_row_num: int,
    data_row_num: int,
    product_attribute: str,
) -> None:
    """指定シートに商品マスタの指定項目を書き込む。"""
    target_column = find_column_by_value(
        worksheet=worksheet,
        search_value=search_value,
        row_number=header_row_num,
    )

    for offset, product_name in enumerate(product_names):
        row_num = data_row_num + offset

        product = product_master.get(product_name)

        if product is None:
            raise ValueError(
                "商品マスタに該当する商品がありません。\n\n"
                f"商品名：{product_name}\n"
                f"対象シート：{worksheet.title}\n"
            )

        worksheet.cell(
            row=row_num,
            column=target_column,
        ).value = getattr(product, product_attribute)

def write_box_table_units(
    worksheet: Worksheet,
    product_master: dict[str, Product],
    product_names: list[str],
    value_column: int,
    start_row: int,
) -> None:
    """箱入り数表の指定列へ商品マスタの1箱当たりの入数を書き込む。"""
    for offset, product_name in enumerate(product_names):
        row_num = start_row + offset

        product = product_master.get(product_name)

        if product is None:
            raise ValueError(
                "商品マスタに該当する商品がありません。\n\n"
                f"商品名：{product_name}\n"
                f"対象シート：{worksheet.title}\n"
            )

        worksheet.cell(
            row=row_num,
            column=value_column,
        ).value = product.units_per_box