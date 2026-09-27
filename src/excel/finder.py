from openpyxl.worksheet.worksheet import Worksheet

from src.utils.converter import normalize_text


def find_row_by_value(worksheet: Worksheet, search_value: str, column_number: int = 1, start_row: int = 1) -> int:
    """指定列から文字列を検索し、該当する行番号を返す。"""
    normalized_search_value = normalize_text(search_value)

    for row_number in range(start_row, worksheet.max_row + 1):
        cell_value = worksheet.cell(row=row_number, column=column_number).value
        if cell_value is None:
            continue
        if normalize_text(cell_value) == normalized_search_value:
            return row_number

    raise ValueError(
        f"「{search_value}」のセルが見つかりませんでした。\n"
        f"対象シート：{worksheet.title}\n"
        f"検索列：{column_number}列目"
    )

def find_column_by_value(worksheet: Worksheet, search_value: str, row_number: int, start_column: int = 1) -> int:
    """指定行から文字列を検索し、該当する列番号を返す。"""
    normalized_search_value = normalize_text(search_value)

    for column_number in range(start_column, worksheet.max_column + 1):
        cell_value = worksheet.cell(row=row_number, column=column_number).value
        if cell_value is None:
            continue
        if normalize_text(cell_value) == normalized_search_value:
            return column_number

    raise ValueError(
        f"「{search_value}」のセルが見つかりませんでした。\n"
        f"対象シート：{worksheet.title}\n"
        f"検索行：{row_number}行目"
    )

def find_branch_columns(worksheet: Worksheet, branch_names: list[str], header_row: int) -> dict[str, int]:
    """
    指定したヘッダー行から営業所名を検索し、営業所名と列番号の対応を返す。
    """
    normalized_branch_names = {
        normalize_text(branch_name): branch_name
        for branch_name in branch_names
    }
    branch_columns: dict[str, int] = {}

    for column_number in range(1, worksheet.max_column + 1):
        cell_value = worksheet.cell(row=header_row, column=column_number).value

        if cell_value is None:
            continue

        normalized_value = normalize_text(cell_value)

        if normalized_value in normalized_branch_names:
            original_branch_name = normalized_branch_names[normalized_value]
            branch_columns[original_branch_name] = column_number

    missing_branches = [branch_name for branch_name in branch_names if branch_name not in branch_columns]

    if missing_branches:
        raise ValueError(
            "集計表に営業所の列が見つかりませんでした。\n\n"
            f"対象シート：{worksheet.title}\n"
            f"検索行：{header_row}行目\n"
            f"見つからない営業所："
            f"{', '.join(missing_branches)}"
        )

    return branch_columns
