from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.constants.excel import MASTER_ID_COLUMN, MASTER_PRODUCT_NAME_COLUMN, MASTER_UNITS_PER_BOX_COLUMN, \
    MASTER_PACKAGING_FLOOR_COLUMN, PRODUCT_MASTER_SHEET_NAME, MASTER_START_ROW, MASTER_ORDER_MULTIPLIER_COLUMN, \
    MASTER_STORE_SALES_UNITS_COLUMN, MASTER_UNIT_PRICE_COLUMN, MASTER_MARGIN_UNIT_PRICE_COLUMN
from src.models.product import Product
from src.utils.converter import normalize_text, normalize_int, normalize_positive_float


def load_product_master(path: Path) -> dict[str, Product]:
    """
    商品マスタシートから商品情報を読み込む。
    """
    workbook = load_workbook(filename=path, data_only=True, read_only=True)

    try:
        if PRODUCT_MASTER_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"商品マスタに必要なシートが見つかりません。\n\n"
                f"必要なシート：{PRODUCT_MASTER_SHEET_NAME}"
            )

        worksheet = workbook[PRODUCT_MASTER_SHEET_NAME]

        products: dict[str, Product] = {}
        product_ids: set[int] = set()

        for row_number in range(MASTER_START_ROW, worksheet.max_row + 1):
            # セルから値を読み込み
            product_id = worksheet.cell(row=row_number, column=MASTER_ID_COLUMN).value
            product_name = worksheet.cell(row=row_number, column=MASTER_PRODUCT_NAME_COLUMN).value
            if product_id is None and product_name is None:
                continue
            units_per_box = worksheet.cell(row=row_number, column=MASTER_UNITS_PER_BOX_COLUMN).value
            packaging_floor = worksheet.cell(row=row_number, column=MASTER_PACKAGING_FLOOR_COLUMN).value
            order_multiplier = worksheet.cell(row=row_number, column=MASTER_ORDER_MULTIPLIER_COLUMN).value
            store_sales_units = worksheet.cell(row=row_number, column=MASTER_STORE_SALES_UNITS_COLUMN).value
            unit_price = worksheet.cell(row=row_number, column=MASTER_UNIT_PRICE_COLUMN).value
            margin_unit_price = worksheet.cell(row=row_number, column=MASTER_MARGIN_UNIT_PRICE_COLUMN).value

            # 値を正規化
            product_id = normalize_int(product_id)
            product_name = normalize_text(product_name)
            units_per_box = normalize_int(units_per_box)
            packaging_floor = normalize_int(packaging_floor)
            order_multiplier = normalize_int(order_multiplier)
            store_sales_units = normalize_int(store_sales_units)
            unit_price = normalize_int(unit_price)
            margin_unit_price = normalize_positive_float(margin_unit_price)

            if not product_name:
                raise ValueError(f"商品マスタの{row_number}行目に商品名がありません。")

            if product_id in product_ids:
                raise ValueError(f"商品マスタに同じ商品IDが重複しています。\n\n商品ID：{product_id}")

            if product_name in products:
                raise ValueError(f"商品マスタに同じ商品名が重複しています。\n\n商品名：{product_name}")

            if packaging_floor not in (1, 2):
                raise ValueError(
                    f"商品マスタのラベル包装階が不正です。\n\n"
                    f"商品名：{product_name}\n"
                    f"値：{packaging_floor}"
                )

            products[product_name] = Product(
                product_id=product_id,
                name=product_name,
                units_per_box=units_per_box,
                packaging_floor=packaging_floor,
                order_multiplier=order_multiplier,
                store_sales_units=store_sales_units,
                unit_price=unit_price,
                margin_unit_price=margin_unit_price,
            )

            product_ids.add(product_id)

        if not products:
            raise ValueError("商品マスタから商品情報を取得できませんでした。")

        return products

    finally:
        workbook.close()


def product_master_to_dataframe(product_master: dict[str, Product]) -> pd.DataFrame:
    return pd.DataFrame({"商品名": list(product_master.keys())})