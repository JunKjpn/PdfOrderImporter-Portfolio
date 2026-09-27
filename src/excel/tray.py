import shutil
from datetime import date, datetime, time
from pathlib import Path

import win32com
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.constants.products import TRAY_PRODUCT_NAME, TRAY_PRODUCT_NAME_A, TRAY_PRODUCT_NAME_B
from src.utils.converter import normalize_amount, normalize_text
from src.constants.excel import TRAY_SHEET_NAME, TRAY_PRODUCT_COLUMN, TRAY_HEADER_ROW, \
    TRAY_TOTAL_NUMBER_LABEL, TRAY_STORE_SALES_UNITS_LABEL, TRAY_TABLE_SHEET_NAME, TRAY_START_ROW
from src.excel.finder import find_row_by_value, find_column_by_value


TRAY_PRODUCT_QUANTITY_COLUMNS = {
    1: 3,  # A列の商品名 → C列へ数量
    4: 6,  # D列の商品名 → F列へ数量
    7: 9,  # G列の商品名 → I列へ数量
}


def load_tray_data(summary_path: Path) -> dict[str, dict[str, int | float]]:
    """
    集計表の「天板換算表」シートから、商品ごとの総数と店売り数を取得する。
    戻り値例:
    {
        "商品C": {
            "total": 120,
            "store_sales": 5,
        },
    }
    """
    workbook = load_workbook(filename=summary_path, data_only=True, read_only=True)

    try:
        if TRAY_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                "集計表に必要なシートが見つかりません。\n\n"
                f"必要なシート：{TRAY_SHEET_NAME}"
            )

        worksheet = workbook[TRAY_SHEET_NAME]

        total_column = find_column_by_value(worksheet=worksheet, search_value=TRAY_TOTAL_NUMBER_LABEL, row_number=TRAY_HEADER_ROW)
        store_sales_column = find_column_by_value(
            worksheet=worksheet, search_value=TRAY_STORE_SALES_UNITS_LABEL, row_number=TRAY_HEADER_ROW
        )
        end_row = find_row_by_value(worksheet=worksheet, search_value="合計", column_number=TRAY_PRODUCT_COLUMN)

        tray_data: dict[str, dict[str, int | float]] = {}
        for row_number in range(TRAY_START_ROW, end_row):
            product_value = worksheet.cell(row=row_number, column=TRAY_PRODUCT_COLUMN).value
            if product_value is None:
                continue

            product_name = normalize_text(product_value)
            if not product_name:
                continue

            total_value = worksheet.cell(row=row_number, column=total_column).value
            store_sales_value = worksheet.cell(row=row_number, column=store_sales_column).value
            tray_data[product_name] = {
                "total": normalize_amount(total_value),
                "store_sales": normalize_amount(store_sales_value)
            }

        if not tray_data:
            raise ValueError(f"天板換算表から商品データを取得できませんでした。\n\n対象シート：{worksheet.title}")

        data_a = tray_data.get(
            TRAY_PRODUCT_NAME_A, {"total": 0, "store_sales": 0}
        )
        data_b = tray_data.get(
            TRAY_PRODUCT_NAME_B, {"total": 0, "store_sales": 0}
        )
        tray_data[normalize_text(TRAY_PRODUCT_NAME)] = {
            "total": data_a["total"] + data_b["total"],
            "store_sales": data_a["store_sales"] + data_b["store_sales"],
        }
        tray_data.pop(normalize_text(TRAY_PRODUCT_NAME_A), None)
        tray_data.pop(normalize_text(TRAY_PRODUCT_NAME_B), None)

        return tray_data

    finally:
        workbook.close()

def find_tray_product_cells(worksheet: Worksheet) -> dict[str, tuple[int, int]]:
    """
    天板並べ早見表の商品名を検索し、商品名と数量転記先セルの対応を返す。
    戻り値例:
    {
        "商品C": (2, 3),
        "商品G": (9, 6),
    }
    """
    product_cells: dict[str, tuple[int, int]] = {}

    for product_column, quantity_column in TRAY_PRODUCT_QUANTITY_COLUMNS.items():
        for row_number in range(1, worksheet.max_row + 1):
            cell_value = worksheet.cell(row=row_number, column=product_column).value
            if not isinstance(cell_value, str):
                continue

            product_name = normalize_text(cell_value)
            if not product_name:
                continue

            product_cells[product_name] = (row_number, quantity_column)

    if not product_cells:
        raise ValueError(
            "天板並べ早見表に商品名が見つかりませんでした。\n\n"
            f"対象シート：{worksheet.title}"
        )

    return product_cells

def transfer_tray(summary_path: Path, tray_path: Path, delivery_date: date, output_dir: Path) -> Path:
    """"
    集計表の天板換算情報をtodayシートへ転記する。
    セル変更時のVBAイベントを動作させるため、書き込みと保存にはExcel本体を使用する。
    """
    # 商品名と転記先セルの位置だけopenpyxlで取得
    tray_data = load_tray_data(summary_path=summary_path)
    read_workbook = load_workbook(filename=tray_path, read_only=True, data_only=False, keep_vba=True)

    try:
        if TRAY_TABLE_SHEET_NAME not in read_workbook.sheetnames:
            raise ValueError(f"天板並べ早見表に転記先シートが見つかりません。\n\n必要なシート：{TRAY_TABLE_SHEET_NAME}")

        read_worksheet = read_workbook[TRAY_TABLE_SHEET_NAME]
        product_cells = find_tray_product_cells(worksheet=read_worksheet)

    finally:
        read_workbook.close()

    # 出力先へテンプレートをコピー
    output_path = (
        output_dir
        / f"{tray_path.stem}_{delivery_date:%Y%m%d}{tray_path.suffix}"
    )
    shutil.copy2(tray_path, output_path)

    # win32で書き込み
    excel = None
    workbook = None

    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.EnableEvents = True

        workbook = excel.Workbooks.Open(str(output_path.resolve()))
        worksheet = workbook.Worksheets(TRAY_TABLE_SHEET_NAME)

        for product_name, values in tray_data.items():
            normalized_product_name = normalize_text(product_name)
            if normalized_product_name not in product_cells:
                print(f"天板並べ早見表に商品がないためスキップ: {product_name}")
                continue

            row_number, quantity_column = product_cells[normalized_product_name]
            total_quantity = values["total"]
            store_sales_quantity = values["store_sales"]

            # COM経由で入力することでWorksheet_Changeを発火させる
            worksheet.Cells(row_number, quantity_column).Value = total_quantity
            worksheet.Cells(row_number + 1, quantity_column).Value = store_sales_quantity

        excel_serial = (datetime.combine(delivery_date, time.min) - datetime(1899, 12, 30)).days
        worksheet.Cells(1, 1).Value = excel_serial
        worksheet.Cells(1, 1).NumberFormat = "yyyy/m/d"
        # 数式や連動処理を念のため再計算
        excel.CalculateFullRebuild()
        workbook.Save()
        return output_path

    except Exception as error:
        raise RuntimeError(
            "天板並べ早見表への転記に失敗しました。\nファイルをExcelで開いている場合は閉じてから、もう一度実行してください。\n\n"
            f"{output_path}"
        ) from error

    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=False)

        if excel is not None:
            excel.EnableEvents = True
            excel.Quit()
