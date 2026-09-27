import shutil
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from src.constants.excel import SUMMARY_SHEET_NAME, BOX_CALCULATION_SHEET_NAME, SUMMARY_PRODUCT_COLUMN, \
    SUMMARY_HEADER_ROW, SUMMARY_START_ROW, DELIVERY_PRODUCT_COLUMN, DELIVERY_QUANTITY_COLUMN, DELIVERY_BOX_COLUMN, \
    DELIVERY_START_ROW
from src.excel.finder import find_branch_columns, find_row_by_value
from src.excel.writer import clear_cell_range
from src.constants.branches import BRANCHES, BRANCH_A
from src.utils.converter import normalize_amount, normalize_text


# =========================================================
# 集計表データ読み込み
# =========================================================
def load_delivery_note_data(summary_path: Path) -> dict[str, dict[str, dict[str, int | float]]]:
    """
    集計表から営業所ごとの商品名・受注数・箱数を取得する。
    戻り値例:
    {
        "営業所A": {
            "商品A": {
                "quantity": 10,
                "box": 2,
            },
        },
    }
    """
    workbook = load_workbook(filename=summary_path, data_only=True, read_only=True)

    try:
        if SUMMARY_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"集計表に必要なシートが見つかりません。\n\n必要なシート：{SUMMARY_SHEET_NAME}"
            )

        if BOX_CALCULATION_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"集計表に必要なシートが見つかりません。\n\n必要なシート：{BOX_CALCULATION_SHEET_NAME}"
            )

        order_worksheet = workbook[SUMMARY_SHEET_NAME]
        box_worksheet = workbook[BOX_CALCULATION_SHEET_NAME]

        order_branch_columns = find_branch_columns(
            worksheet=order_worksheet,
            branch_names=BRANCHES,
            header_row=SUMMARY_HEADER_ROW,
        )

        box_branch_columns = find_branch_columns(
            worksheet=box_worksheet,
            branch_names=BRANCHES,
            header_row=SUMMARY_HEADER_ROW,
        )

        end_row = find_row_by_value(
            worksheet=box_worksheet,
            search_value="合計",
            column_number=SUMMARY_PRODUCT_COLUMN,
            start_row=SUMMARY_START_ROW,
        )

        delivery_note_data: dict[str, dict[str, dict[str, int | float]]] = {
            branch_name: {}
            for branch_name in BRANCHES
        }

        for row_number in range(SUMMARY_START_ROW, end_row):
            product_value = box_worksheet.cell(row=row_number, column=SUMMARY_PRODUCT_COLUMN).value
            if product_value is None:
                continue

            product_name = normalize_text(product_value)
            if not product_name:
                continue

            for branch_name in BRANCHES:
                box_column = box_branch_columns.get(normalize_text(branch_name))
                if box_column is None:
                    continue
                box_value = box_worksheet.cell(row=row_number, column=box_column).value

                quantity_column = order_branch_columns.get(normalize_text(branch_name))
                if quantity_column is None:
                    continue
                quantity_value = order_worksheet.cell(row=row_number, column=quantity_column).value

                delivery_note_data[branch_name][product_name] = {
                    "quantity": normalize_amount(quantity_value),
                    "box": normalize_amount(box_value),
                }

        return delivery_note_data

    finally:
        workbook.close()


# =========================================================
# 転記処理
# =========================================================
def transfer_delivery_note(summary_path: Path, delivery_note_path: Path, delivery_date: date, output_dir: Path) -> Path:
    """
    集計表から営業所ごとの受注数・箱数を納品書へ転記する。
    """
    delivery_note_data = load_delivery_note_data(summary_path=summary_path)
    output_path = output_dir / f"{delivery_note_path.stem}_{delivery_date:%Y%m%d}{delivery_note_path.suffix}"
    shutil.copy2(delivery_note_path, output_path)
    workbook = load_workbook(filename=output_path)

    try:
        delivery_sheet_names = {normalize_text(sheet_name): sheet_name for sheet_name in workbook.sheetnames}

        for branch_name, product_data in delivery_note_data.items():
            normalized_branch_name = normalize_text(branch_name)
            if normalized_branch_name not in delivery_sheet_names:
                raise ValueError(f"納品書に営業所シートが見つかりません。\n\n営業所：{branch_name}")

            sheet_name = delivery_sheet_names[normalized_branch_name]
            worksheet = workbook[sheet_name]

            total_row = find_row_by_value(
                worksheet=worksheet, search_value="合計", column_number=DELIVERY_PRODUCT_COLUMN,
                start_row=DELIVERY_START_ROW,
            )

            clear_cell_range(
                worksheet=worksheet,
                start_row=DELIVERY_START_ROW,
                end_row=total_row - 1,
                start_column=DELIVERY_PRODUCT_COLUMN,
                end_column=DELIVERY_BOX_COLUMN,
            )

            target_row = DELIVERY_START_ROW

            for product_name, values in product_data.items():
                worksheet.cell(
                    row=target_row, column=DELIVERY_PRODUCT_COLUMN, value=product_name,
                )

                worksheet.cell(
                    row=target_row, column=DELIVERY_QUANTITY_COLUMN, value=values["quantity"]
                )

                worksheet.cell(row=target_row, column=DELIVERY_BOX_COLUMN, value=values["box"])

                target_row += 1

        south_osaka_sheet_name = delivery_sheet_names[normalize_text(BRANCH_A)]
        south_osaka_worksheet = workbook[south_osaka_sheet_name]
        cell = south_osaka_worksheet.cell(row=1, column=2)
        cell.value = delivery_date
        cell.number_format = 'yyyy"年"m"月"d"日"(aaa)'
        workbook.save(output_path)

        return output_path

    except PermissionError as error:
        raise PermissionError(
            "納品書を保存できませんでした。\nExcelで開いている場合は閉じてから、もう一度実行してください。\n\n"
            f"{output_path}"
        ) from error

    finally:
        workbook.close()
